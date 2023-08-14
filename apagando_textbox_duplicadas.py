from openpyxl import load_workbook

wb = load_workbook(filename = 'Planilha com Problema - REUNIÃO DIÁRIA - CÉLULA 14.xlsx')
ws = wb.active  # considera a planilha ativa

for obj in ws.shapes:
    if "TextBox 164" in obj.name or "TextBox 176" in obj.name:
        ws.shapes.remove(obj)

wb.save('Planilha consertada - REUNIÃO DIÁRIA - CÉLULA 14.xlsx')